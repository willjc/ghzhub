#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
修复2人家庭合同tenant_id映射错误 (优化版: 批量查询)

问题: migrate_data.py step6中, 同一资格证号的两人(主申请人+配偶),
      配偶的id_mapper.register('qualification', qual_no, new_id) 覆盖了主申请人,
      导致合同被错误关联到配偶的user_id。

修复: 通过身份证号精确匹配, 将合同/账单/入住/退租等表的tenant_id从配偶改为主申请人。
"""

import os
import sys
import datetime
import openpyxl
import pymysql
from collections import defaultdict

# ============================================================
# 配置
# ============================================================

DB_CONFIG = {
    'host': '36.133.39.148',
    'port': 33061,
    'user': 'ghzuser',
    'password': '!aNJ-7vTE+Tsddce',
    'database': 'ghz',
    'charset': 'utf8mb4',
    'cursorclass': pymysql.cursors.DictCursor,
}

DATA_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), '系统数据')
PERSON_FILE = os.path.join(DATA_DIR, '人员.xlsx')


def log(msg, level='INFO'):
    timestamp = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    print(f"[{timestamp}] [{level}] {msg}")


def safe_str(val):
    if val is None or str(val).strip() == '' or str(val).strip() == 'None':
        return None
    return str(val).strip()


# ============================================================
# Step 1: 从Excel构建2人家庭映射
# ============================================================

def build_family_mapping():
    log("=" * 60)
    log("Step 1: 从老系统Excel构建家庭映射")
    log("=" * 60)

    wb = openpyxl.load_workbook(PERSON_FILE, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    if not rows:
        log("人员.xlsx为空!", 'ERROR')
        return {}

    headers = [str(h).strip() if h else f'col_{i}' for i, h in enumerate(rows[0])]

    # 按资格证号分组
    families = defaultdict(list)
    for row in rows[1:]:
        d = {headers[i]: row[i] if i < len(row) else None for i in range(len(headers))}
        qual_no = safe_str(d.get('资格证号（一个家庭一个号）'))
        id_card = safe_str(d.get('申请人证件号'))
        name = safe_str(d.get('申请人姓名'))
        apply_type = safe_str(d.get('申请人类型'))

        if not qual_no or not id_card:
            continue

        families[qual_no].append({
            'name': name,
            'id_card': id_card,
            'apply_type': apply_type,
        })

    # 筛选2人家庭并识别主申请人/配偶
    two_person_families = {}
    for qual_no, members in families.items():
        if len(members) != 2:
            continue

        principal = None
        spouse = None
        for m in members:
            t = m['apply_type']
            if t and '配偶' in t:
                spouse = m
            elif t and '主申请人' in t:
                principal = m

        if not principal and not spouse:
            principal, spouse = members[0], members[1]
        elif not principal:
            principal = [m for m in members if m != spouse][0]
        elif not spouse:
            spouse = [m for m in members if m != principal][0]

        two_person_families[qual_no] = {
            'principal': principal,
            'spouse': spouse,
        }

    log(f"  总家庭数: {len(families)}, 2人家庭数: {len(two_person_families)}")
    return two_person_families


# ============================================================
# Step 2-5: 批量排查并修复
# ============================================================

def check_and_fix(families):
    log("=" * 60)
    log("Step 2: 连接数据库, 批量查询user_id")
    log("=" * 60)

    conn = pymysql.connect(**DB_CONFIG)
    cursor = conn.cursor()
    log("  数据库连接成功")

    # 1) 批量加载所有迁移用户 (id_card -> user_id)
    log("  加载所有迁移用户...")
    cursor.execute("SELECT user_id, real_name, id_card, phone FROM hz_user WHERE create_by = 'migration' AND del_flag = '0'")
    all_users = cursor.fetchall()
    idcard_to_user = {u['id_card']: u for u in all_users}
    log(f"  加载了 {len(all_users)} 个迁移用户")

    # 2) 批量加载所有迁移合同
    log("  加载所有迁移合同...")
    cursor.execute("""
        SELECT contract_id, contract_no, tenant_id, tenant_name, tenant_id_card, contract_status
        FROM hz_contract WHERE create_by = 'migration' AND del_flag = '0'
    """)
    all_contracts = cursor.fetchall()
    # 按tenant_id索引
    contracts_by_tenant = defaultdict(list)
    for c in all_contracts:
        if c['tenant_id']:
            contracts_by_tenant[c['tenant_id']].append(c)
    log(f"  加载了 {len(all_contracts)} 条迁移合同")

    # 3) 遍历2人家庭, 找出错误映射
    log("")
    log("=" * 60)
    log("Step 3: 排查错误映射")
    log("=" * 60)

    wrong_contracts = []  # (contract_id, old_tenant_id, new_tenant_id, info)
    families_checked = 0
    families_no_user = 0
    families_wrong = 0

    for qual_no, fam in families.items():
        p_idcard = fam['principal']['id_card']
        s_idcard = fam['spouse']['id_card']

        p_user = idcard_to_user.get(p_idcard)
        s_user = idcard_to_user.get(s_idcard)

        if not p_user or not s_user:
            families_no_user += 1
            continue

        families_checked += 1
        principal_uid = p_user['user_id']
        spouse_uid = s_user['user_id']

        # 配偶名下的合同就是被错误关联的
        spouse_contracts = contracts_by_tenant.get(spouse_uid, [])
        if spouse_contracts:
            families_wrong += 1
            for c in spouse_contracts:
                wrong_contracts.append({
                    'contract_id': c['contract_id'],
                    'contract_no': c['contract_no'],
                    'old_tenant_id': spouse_uid,
                    'new_tenant_id': principal_uid,
                    'principal_name': fam['principal']['name'],
                    'spouse_name': fam['spouse']['name'],
                    'qual_no': qual_no,
                    'contract_status': c['contract_status'],
                })

    log(f"  检查2人家庭: {families_checked} (未找到user: {families_no_user})")
    log(f"  存在错误映射的家庭: {families_wrong}")
    log(f"  需修正的合同: {len(wrong_contracts)}")

    if not wrong_contracts:
        log("  没有需要修正的合同，退出")
        cursor.close()
        conn.close()
        return

    # 打印每条错误合同
    log("")
    log("受影响的合同明细:")
    for i, c in enumerate(wrong_contracts):
        log(f"  {i+1}. 合同ID={c['contract_id']}, 编号={c['contract_no']}, "
            f"状态={c['contract_status']}, 资格证号={c['qual_no']}, "
            f"错误→{c['spouse_name']}(uid={c['old_tenant_id']}), "
            f"应为→{c['principal_name']}(uid={c['new_tenant_id']})")

    # 4) 批量查找关联表受影响记录
    contract_ids = [c['contract_id'] for c in wrong_contracts]
    # 建立 contract_id -> (old_tid, new_tid) 映射
    fix_map = {}
    for c in wrong_contracts:
        fix_map[c['contract_id']] = (c['old_tenant_id'], c['new_tenant_id'])

    # 查账单
    if contract_ids:
        placeholders = ','.join(['%s'] * len(contract_ids))
        cursor.execute(f"SELECT bill_id, contract_id, tenant_id FROM hz_bill WHERE contract_id IN ({placeholders}) AND create_by = 'migration'", contract_ids)
        all_bills = cursor.fetchall()
        wrong_bills = [b for b in all_bills if b['contract_id'] in fix_map and b['tenant_id'] == fix_map[b['contract_id']][0]]

        cursor.execute(f"SELECT record_id, contract_id, tenant_id FROM hz_checkin_record WHERE contract_id IN ({placeholders}) AND create_by = 'migration'", contract_ids)
        all_checkins = cursor.fetchall()
        wrong_checkins = [r for r in all_checkins if r['contract_id'] in fix_map and r['tenant_id'] == fix_map[r['contract_id']][0]]

        cursor.execute(f"SELECT record_id, contract_id, tenant_id FROM hz_checkout_record WHERE contract_id IN ({placeholders}) AND create_by = 'migration'", contract_ids)
        all_checkouts = cursor.fetchall()
        wrong_checkouts = [r for r in all_checkouts if r['contract_id'] in fix_map and r['tenant_id'] == fix_map[r['contract_id']][0]]

        cursor.execute(f"SELECT apply_id, contract_id, tenant_id FROM hz_checkout_apply WHERE contract_id IN ({placeholders}) AND create_by = 'migration'", contract_ids)
        all_applies = cursor.fetchall()
        wrong_applies = [a for a in all_applies if a['contract_id'] in fix_map and a['tenant_id'] == fix_map[a['contract_id']][0]]

        # 缴费记录 (通过bill_id关联)
        bill_ids = [b['bill_id'] for b in wrong_bills]
        wrong_payments = []
        if bill_ids:
            bp = ','.join(['%s'] * len(bill_ids))
            cursor.execute(f"SELECT payment_id, bill_id, tenant_id FROM hz_payment WHERE bill_id IN ({bp}) AND create_by = 'migration'", bill_ids)
            all_payments = cursor.fetchall()
            # 建立bill_id -> (old, new)
            bill_fix = {}
            for b in wrong_bills:
                old_tid, new_tid = fix_map[b['contract_id']]
                bill_fix[b['bill_id']] = (old_tid, new_tid)
            wrong_payments = [p for p in all_payments if p['bill_id'] in bill_fix and p['tenant_id'] == bill_fix[p['bill_id']][0]]

        # 签署记录
        cursor.execute(f"SELECT signature_id, contract_id, signer_id FROM hz_contract_signature WHERE contract_id IN ({placeholders}) AND create_by = 'migration'", contract_ids)
        all_sigs = cursor.fetchall()
        wrong_sigs = [s for s in all_sigs if s['contract_id'] in fix_map and s['signer_id'] == fix_map[s['contract_id']][0]]
    else:
        wrong_bills = wrong_checkins = wrong_checkouts = wrong_applies = wrong_payments = wrong_sigs = []

    log("")
    log("=" * 60)
    log("排查报告")
    log("=" * 60)
    log(f"  错误映射家庭数: {families_wrong}")
    log(f"  需修正合同: {len(wrong_contracts)}")
    log(f"  需修正账单: {len(wrong_bills)}")
    log(f"  需修正入住记录: {len(wrong_checkins)}")
    log(f"  需修正退租记录: {len(wrong_checkouts)}")
    log(f"  需修正退租申请: {len(wrong_applies)}")
    log(f"  需修正缴费记录: {len(wrong_payments)}")
    log(f"  需修正签署记录: {len(wrong_sigs)}")

    # ========================================
    # Step 4: 执行修正
    # ========================================
    log("")
    log("=" * 60)
    log("Step 4: 执行修正")
    log("=" * 60)

    try:
        # 修正合同
        fc = 0
        for c in wrong_contracts:
            cursor.execute("""
                UPDATE hz_contract
                SET tenant_id = %s,
                    tenant_name = (SELECT real_name FROM hz_user WHERE user_id = %s),
                    tenant_id_card = (SELECT id_card FROM hz_user WHERE user_id = %s),
                    tenant_phone = (SELECT phone FROM hz_user WHERE user_id = %s)
                WHERE contract_id = %s AND tenant_id = %s AND create_by = 'migration'
            """, (c['new_tenant_id'], c['new_tenant_id'], c['new_tenant_id'], c['new_tenant_id'],
                  c['contract_id'], c['old_tenant_id']))
            fc += cursor.rowcount
        log(f"  修正合同: {fc}")

        # 修正账单
        fb = 0
        for b in wrong_bills:
            old_tid, new_tid = fix_map[b['contract_id']]
            cursor.execute("""
                UPDATE hz_bill SET tenant_id = %s,
                    tenant_name = (SELECT real_name FROM hz_user WHERE user_id = %s)
                WHERE bill_id = %s AND tenant_id = %s
            """, (new_tid, new_tid, b['bill_id'], old_tid))
            fb += cursor.rowcount
        log(f"  修正账单: {fb}")

        # 修正入住记录
        fi = 0
        for r in wrong_checkins:
            old_tid, new_tid = fix_map[r['contract_id']]
            cursor.execute("UPDATE hz_checkin_record SET tenant_id = %s WHERE record_id = %s AND tenant_id = %s",
                           (new_tid, r['record_id'], old_tid))
            fi += cursor.rowcount
        log(f"  修正入住记录: {fi}")

        # 修正退租记录
        fo = 0
        for r in wrong_checkouts:
            old_tid, new_tid = fix_map[r['contract_id']]
            cursor.execute("UPDATE hz_checkout_record SET tenant_id = %s WHERE record_id = %s AND tenant_id = %s",
                           (new_tid, r['record_id'], old_tid))
            fo += cursor.rowcount
        log(f"  修正退租记录: {fo}")

        # 修正退租申请
        fa = 0
        for a in wrong_applies:
            old_tid, new_tid = fix_map[a['contract_id']]
            cursor.execute("UPDATE hz_checkout_apply SET tenant_id = %s WHERE apply_id = %s AND tenant_id = %s",
                           (new_tid, a['apply_id'], old_tid))
            fa += cursor.rowcount
        log(f"  修正退租申请: {fa}")

        # 修正缴费记录
        fp = 0
        if wrong_payments:
            bill_fix = {}
            for b in wrong_bills:
                old_tid, new_tid = fix_map[b['contract_id']]
                bill_fix[b['bill_id']] = (old_tid, new_tid)
            for p in wrong_payments:
                old_tid, new_tid = bill_fix[p['bill_id']]
                cursor.execute("""
                    UPDATE hz_payment SET tenant_id = %s,
                        tenant_name = (SELECT real_name FROM hz_user WHERE user_id = %s)
                    WHERE payment_id = %s AND tenant_id = %s
                """, (new_tid, new_tid, p['payment_id'], old_tid))
                fp += cursor.rowcount
        log(f"  修正缴费记录: {fp}")

        # 修正签署记录
        fs = 0
        for s in wrong_sigs:
            old_tid, new_tid = fix_map[s['contract_id']]
            cursor.execute("""
                UPDATE hz_contract_signature SET signer_id = %s,
                    signer_name = (SELECT real_name FROM hz_user WHERE user_id = %s)
                WHERE signature_id = %s AND signer_id = %s
            """, (new_tid, new_tid, s['signature_id'], old_tid))
            fs += cursor.rowcount
        log(f"  修正签署记录: {fs}")

        conn.commit()
        log("  事务已提交!")

    except Exception as e:
        conn.rollback()
        log(f"  修正失败，已回滚: {e}", 'ERROR')
        import traceback
        traceback.print_exc()
        cursor.close()
        conn.close()
        return

    # ========================================
    # Step 5: 验证
    # ========================================
    log("")
    log("=" * 60)
    log("Step 5: 验证修正结果")
    log("=" * 60)

    # 验证潘静
    cursor.execute("SELECT user_id, real_name FROM hz_user WHERE id_card = '410882198909261044' AND del_flag = '0'")
    panjing = cursor.fetchone()
    if panjing:
        cursor.execute("""
            SELECT contract_id, contract_no, tenant_id, tenant_name, contract_status, start_date, end_date, house_code
            FROM hz_contract WHERE tenant_id = %s AND del_flag = '0'
        """, (panjing['user_id'],))
        pj_contracts = cursor.fetchall()
        log(f"  潘静(user_id={panjing['user_id']}): {len(pj_contracts)} 条合同")
        for c in pj_contracts:
            log(f"    合同ID={c['contract_id']}, 编号={c['contract_no']}, 状态={c['contract_status']}, "
                f"房间={c['house_code']}, 期间={c['start_date']}~{c['end_date']}")
    else:
        log("  潘静(410882198909261044): 未在hz_user中找到!", 'WARN')

    # 验证所有修正合同
    log("")
    all_ok = True
    for c in wrong_contracts:
        cursor.execute("SELECT tenant_id, tenant_name FROM hz_contract WHERE contract_id = %s", (c['contract_id'],))
        r = cursor.fetchone()
        if r and r['tenant_id'] != c['new_tenant_id']:
            log(f"  ✗ 合同{c['contract_id']}: tenant_id={r['tenant_id']} 期望={c['new_tenant_id']}", 'ERROR')
            all_ok = False

    if all_ok:
        log(f"  ✓ 所有 {len(wrong_contracts)} 条合同tenant_id已修正为主申请人")

    # 验证关联表无残留
    for c in wrong_contracts:
        cid = c['contract_id']
        old_tid = c['old_tenant_id']
        for tbl, id_col in [('hz_bill', 'bill_id'), ('hz_checkin_record', 'record_id'),
                             ('hz_checkout_record', 'record_id'), ('hz_checkout_apply', 'apply_id')]:
            cursor.execute(f"SELECT COUNT(*) as cnt FROM {tbl} WHERE contract_id = %s AND tenant_id = %s AND create_by = 'migration'",
                           (cid, old_tid))
            r = cursor.fetchone()
            if r['cnt'] > 0:
                log(f"  ✗ 合同{cid}: {tbl}残留 {r['cnt']} 条旧tenant_id", 'WARN')

    log("")
    log("=" * 60)
    log("修正报告总结")
    log("=" * 60)
    log(f"  受影响2人家庭数: {families_wrong}")
    log(f"  修正合同: {fc}")
    log(f"  修正账单: {fb}")
    log(f"  修正入住记录: {fi}")
    log(f"  修正退租记录: {fo}")
    log(f"  修正退租申请: {fa}")
    log(f"  修正缴费记录: {fp}")
    log(f"  修正签署记录: {fs}")

    cursor.close()
    conn.close()
    log("  数据库连接已关闭")


def main():
    log("修复2人家庭合同tenant_id映射错误")
    families = build_family_mapping()
    if not families:
        log("没有找到2人家庭，退出")
        return
    check_and_fix(families)
    log("全部完成!")


if __name__ == '__main__':
    main()
